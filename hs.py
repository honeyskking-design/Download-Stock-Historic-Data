import re
import time
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError


# ============================================================
# CONFIGURATION
# ============================================================

BASE_FOLDER = Path(
    r"D:\Documents(H)\Company historic data - Code"
)

INPUT_FILE = BASE_FOLDER / "CompanyTicker.xlsx"

OUTPUT_FOLDER = BASE_FOLDER / "Company Historic Data"

NSE_URL = "https://www.nsehistoricaldata.co.in/"
BSE_URL = "https://bse.nsehistoricaldata.co.in/"

FROM_DATE = "2015-04-01"
TO_DATE = "2026-03-31"


# ============================================================
# DOWNLOAD TIMEOUT
# ============================================================
# Maximum time to wait for the Excel download to START.
#
# 120000 milliseconds = 120 seconds = 2 minutes.
#
# IMPORTANT:
# This is NOT a mandatory 2-minute wait.
# If the download starts in 20 seconds, the script continues
# after 20 seconds.
# ============================================================

DOWNLOAD_TIMEOUT = 120000


# Small pause between companies
DELAY_BETWEEN_COMPANIES = 2


# Run browser completely in the background
HEADLESS = True


# ============================================================
# OUTPUT FILES
# ============================================================

MASTER_FILE = OUTPUT_FOLDER / "Master_Historical_Data.xlsx"

LOG_FILE = OUTPUT_FOLDER / "Download_Log.xlsx"


# ============================================================
# READ INPUT TICKERS
# ============================================================

def read_tickers():

    print("\nReading input Excel...")

    try:
        df = pd.read_excel(
            INPUT_FILE,
            header=0
        )

    except Exception as e:

        raise RuntimeError(
            f"Could not open input file '{INPUT_FILE}'.\n"
            f"Error: {e}"
        )

    if df.empty:

        raise RuntimeError(
            "The input Excel file is empty."
        )

    # --------------------------------------------------------
    # Current setup:
    # The script reads tickers from Column A.
    # --------------------------------------------------------

    ticker_series = df.iloc[:, 0]

    tickers = []

    for value in ticker_series:

        if pd.isna(value):
            continue

        ticker = str(value).strip()

        if not ticker:
            continue

        tickers.append(ticker)

    # Remove duplicates while preserving order
    tickers = list(
        dict.fromkeys(tickers)
    )

    print(
        f"Found {len(tickers)} unique ticker(s)."
    )

    # --------------------------------------------------------
    # Prevent the previous KeyError: 'Status'
    # --------------------------------------------------------

    if len(tickers) == 0:

        print("\nWARNING:")
        print(
            "No tickers were found in Column A "
            "of CompanyTicker.xlsx."
        )

        print(
            "\nThe program cannot continue because "
            "there are no companies to download."
        )

        return []

    return tickers


# ============================================================
# CLEAN FILENAMES
# ============================================================

def clean_filename(name):

    name = str(name).strip()

    # Windows-invalid filename characters
    name = re.sub(
        r'[<>:"/\\|?*]',
        "_",
        name
    )

    return name


# ============================================================
# CLEAN EXCEL SHEET NAME
# ============================================================

def clean_sheet_name(name):

    name = str(name).strip()

    # Excel does not allow:
    # \ / ? * [ ]
    name = re.sub(
        r'[\\/*?:\[\]]',
        "_",
        name
    )

    if not name:

        name = "Company"

    # Excel sheet names max 31 characters
    return name[:31]


# ============================================================
# CREATE UNIQUE SHEET NAME
# ============================================================

def get_unique_sheet_name(
    workbook,
    requested_name
):

    base_name = clean_sheet_name(
        requested_name
    )

    if base_name not in workbook.sheetnames:

        return base_name

    counter = 2

    while True:

        suffix = f"_{counter}"

        # Maximum Excel sheet name length = 31
        available_length = (
            31 - len(suffix)
        )

        candidate = (
            base_name[:available_length]
            + suffix
        )

        if candidate not in workbook.sheetnames:

            return candidate

        counter += 1


# ============================================================
# INITIALIZE MASTER WORKBOOK
# ============================================================

def initialize_master_workbook():

    if MASTER_FILE.exists():

        return

    print(
        "\nCreating master workbook..."
    )

    workbook = Workbook()

    # Remove default Sheet
    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )

    # Create README sheet
    info = workbook.create_sheet(
        "README"
    )

    info["A1"] = (
        "Historical Stock Data Master Workbook"
    )

    info["A3"] = "Date From"
    info["B3"] = FROM_DATE

    info["A4"] = "Date To"
    info["B4"] = TO_DATE

    info["A5"] = "Source"

    info["B5"] = (
        "NSE Historical Data / "
        "BSE Historical Data"
    )

    info["A7"] = (
        "Each company is stored on "
        "a separate worksheet."
    )

    info["A8"] = (
        "Individual downloaded files are stored "
        "in the same Company Historic Data folder."
    )

    workbook.save(
        MASTER_FILE
    )


# ============================================================
# ADD COMPANY DATA TO MASTER WORKBOOK
# ============================================================

def add_to_master_workbook(
    excel_file,
    ticker,
    exchange
):

    print(
        f"Adding {ticker} ({exchange}) "
        f"to master workbook..."
    )

    try:

        # Read downloaded Excel
        df = pd.read_excel(
            excel_file
        )

        if df.empty:

            raise RuntimeError(
                "Downloaded Excel file "
                "contains no data."
            )

        # Load master workbook
        workbook = load_workbook(
            MASTER_FILE
        )

        # Requested sheet name
        requested_sheet_name = (
            f"{ticker}_{exchange}"
        )

        sheet_name = (
            get_unique_sheet_name(
                workbook,
                requested_sheet_name
            )
        )

        worksheet = workbook.create_sheet(
            sheet_name
        )

        # ----------------------------------------------------
        # Write dataframe headers
        # ----------------------------------------------------

        for column_number, column_name in enumerate(
            df.columns,
            start=1
        ):

            worksheet.cell(
                row=1,
                column=column_number,
                value=str(column_name)
            )

        # ----------------------------------------------------
        # Write dataframe data
        # ----------------------------------------------------

        for row_number, row in enumerate(
            df.itertuples(index=False),
            start=2
        ):

            for column_number, value in enumerate(
                row,
                start=1
            ):

                # Convert NaN to blank
                if pd.isna(value):

                    value = None

                worksheet.cell(
                    row=row_number,
                    column=column_number,
                    value=value
                )

        # ----------------------------------------------------
        # Basic formatting
        # ----------------------------------------------------

        worksheet.freeze_panes = "A2"

        # Reasonable column widths
        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = (
                column_cells[0].column_letter
            )

            for cell in column_cells:

                try:

                    value_length = len(
                        str(cell.value)
                    )

                    if value_length > max_length:

                        max_length = (
                            value_length
                        )

                except Exception:

                    pass

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 10),
                30
            )

        workbook.save(
            MASTER_FILE
        )

        print(
            f"Master workbook updated: "
            f"{sheet_name}"
        )

        return True

    except Exception as e:

        print(
            f"ERROR adding {ticker} to master: {e}"
        )

        return False


# ============================================================
# SAVE DOWNLOAD LOG
# ============================================================

def save_log(results):

    df = pd.DataFrame(
        results
    )

    df.to_excel(
        LOG_FILE,
        index=False
    )

    print(
        f"\nDownload log saved: "
        f"{LOG_FILE.resolve()}"
    )


# ============================================================
# WAIT FOR PAGE
# ============================================================

def wait_for_page(page):

    try:

        page.wait_for_load_state(
            "domcontentloaded",
            timeout=30000
        )

    except Exception:

        pass

    try:

        page.wait_for_load_state(
            "networkidle",
            timeout=15000
        )

    except Exception:

        # Not fatal.
        # Some pages continue background requests.
        pass


# ============================================================
# FIND SYMBOL INPUT
# ============================================================

def find_symbol_input(page):

    selectors = [

        'input[placeholder*="stock name or symbol"]',

        'input[placeholder*="stock"]',

        'input[placeholder*="symbol"]',

        'input[type="text"]'
    ]

    for selector in selectors:

        locator = page.locator(
            selector
        )

        try:

            count = locator.count()

            if count > 0:

                for i in range(count):

                    candidate = locator.nth(i)

                    if candidate.is_visible():

                        return candidate

        except Exception:

            continue

    raise RuntimeError(
        "Could not find the stock symbol input."
    )


# ============================================================
# ENTER TICKER
# ============================================================

def enter_ticker(
    page,
    ticker
):

    symbol_input = find_symbol_input(
        page
    )

    symbol_input.fill("")

    symbol_input.fill(
        ticker
    )

    # Allow autocomplete to populate
    page.wait_for_timeout(
        1500
    )


# ============================================================
# SELECT AUTOCOMPLETE RESULT
# ============================================================

def select_matching_suggestion(
    page,
    ticker
):

    ticker_upper = (
        ticker.strip().upper()
    )

    # Give autocomplete a little additional time
    page.wait_for_timeout(
        1000
    )

    # Search visible text
    candidates = page.locator(
        f"text={ticker}"
    )

    try:

        count = candidates.count()

    except Exception:

        count = 0

    if count == 0:

        return False

    for i in range(
        min(count, 20)
    ):

        candidate = candidates.nth(i)

        try:

            if not candidate.is_visible():

                continue

            text = (
                candidate.inner_text()
                .strip()
            )

            if ticker_upper in text.upper():

                candidate.click()

                page.wait_for_timeout(
                    500
                )

                return True

        except Exception:

            continue

    return False


# ============================================================
# SET DATES
# ============================================================

def set_dates(page):

    date_inputs = page.locator(
        'input[type="date"]'
    )

    count = date_inputs.count()

    if count < 2:

        raise RuntimeError(
            f"Expected two date fields, "
            f"but found {count}."
        )

    # From date
    date_inputs.nth(0).fill(
        FROM_DATE
    )

    # To date
    date_inputs.nth(1).fill(
        TO_DATE
    )


# ============================================================
# FIND DOWNLOAD BUTTON
# ============================================================

def find_download_button(page):

    # Preferred: exact accessible button
    button = page.get_by_role(
        "button",
        name=re.compile(
            "download historical data",
            re.IGNORECASE
        )
    )

    try:

        if button.count() > 0:

            for i in range(
                button.count()
            ):

                candidate = button.nth(i)

                if candidate.is_visible():

                    return candidate

    except Exception:

        pass

    # Fallback
    button = page.locator(
        'button:has-text("Download historical data")'
    )

    if button.count() > 0:

        return button.first

    raise RuntimeError(
        "Could not find Download Historical Data button."
    )


# ============================================================
# DOWNLOAD EXCEL
# ============================================================

def download_excel(page):

    button = find_download_button(
        page
    )

    print(
        "\nWaiting for website to generate Excel..."
    )

    print(
        "Maximum download wait: "
        f"{DOWNLOAD_TIMEOUT // 60000} minutes"
    )

    try:

        with page.expect_download(
            timeout=DOWNLOAD_TIMEOUT
        ) as download_info:

            button.click()

        download = (
            download_info.value
        )

        return download

    except PlaywrightTimeoutError:

        raise RuntimeError(
            "Download did not start within "
            f"{DOWNLOAD_TIMEOUT // 1000} seconds."
        )


# ============================================================
# SAVE INDIVIDUAL DOWNLOAD
# ============================================================

def save_individual_download(
    download,
    ticker,
    exchange
):

    OUTPUT_FOLDER.mkdir(
        parents=True,
        exist_ok=True
    )

    safe_ticker = clean_filename(
        ticker
    )

    filename = (
        f"{safe_ticker}_"
        f"{exchange}_"
        f"2015-04-01_to_2026-03-31.xlsx"
    )

    output_path = (
        OUTPUT_FOLDER / filename
    )

    download.save_as(
        str(output_path)
    )

    return output_path


# ============================================================
# TRY NSE
# ============================================================

def try_nse(
    page,
    ticker
):

    print("\n----------------------------------------")
    print(
        f"NSE CHECK: {ticker}"
    )
    print("----------------------------------------")

    try:

        page.goto(
            NSE_URL,
            wait_until="domcontentloaded",
            timeout=60000
        )

        wait_for_page(
            page
        )

        enter_ticker(
            page,
            ticker
        )

        matched = (
            select_matching_suggestion(
                page,
                ticker
            )
        )

        # ----------------------------------------------------
        # IMPORTANT:
        # If no NSE autocomplete match is found,
        # immediately return to main() so that BSE is tried.
        #
        # We DO NOT attempt the NSE download in this case.
        # ----------------------------------------------------

        if not matched:

            print(
                f"No NSE autocomplete match "
                f"found for {ticker}."
            )

            print(
                "Moving to BSE immediately."
            )

            return (
                False,
                None,
                f"No NSE autocomplete match found "
                f"for {ticker}"
            )

        print(
            f"NSE match found for {ticker}"
        )

        # ----------------------------------------------------
        # NSE match confirmed
        # ----------------------------------------------------

        set_dates(
            page
        )

        download = download_excel(
            page
        )

        output_path = (
            save_individual_download(
                download,
                ticker,
                "NSE"
            )
        )

        print(
            f"NSE DOWNLOAD SUCCESS: "
            f"{output_path.name}"
        )

        return (
            True,
            output_path,
            None
        )

    except Exception as e:

        error = str(e)

        print(
            f"NSE FAILED: {error}"
        )

        return (
            False,
            None,
            error
        )


# ============================================================
# TRY BSE
# ============================================================

def try_bse(
    page,
    ticker
):

    print("\n----------------------------------------")
    print(
        f"BSE CHECK: {ticker}"
    )
    print("----------------------------------------")

    try:

        page.goto(
            BSE_URL,
            wait_until="domcontentloaded",
            timeout=60000
        )

        wait_for_page(
            page
        )

        enter_ticker(
            page,
            ticker
        )

        matched = (
            select_matching_suggestion(
                page,
                ticker
            )
        )

        # ----------------------------------------------------
        # If BSE autocomplete doesn't find the ticker,
        # don't unnecessarily wait for a download.
        # ----------------------------------------------------

        if not matched:

            print(
                f"No BSE autocomplete match "
                f"found for {ticker}."
            )

            return (
                False,
                None,
                f"No BSE autocomplete match found "
                f"for {ticker}"
            )

        print(
            f"BSE match found for {ticker}"
        )

        # ----------------------------------------------------
        # BSE match confirmed
        # ----------------------------------------------------

        set_dates(
            page
        )

        download = download_excel(
            page
        )

        output_path = (
            save_individual_download(
                download,
                ticker,
                "BSE"
            )
        )

        print(
            f"BSE DOWNLOAD SUCCESS: "
            f"{output_path.name}"
        )

        return (
            True,
            output_path,
            None
        )

    except Exception as e:

        error = str(e)

        print(
            f"BSE FAILED: {error}"
        )

        return (
            False,
            None,
            error
        )


# ============================================================
# MAIN
# ============================================================

def main():

    print("\n")
    print("=" * 75)

    print(
        "       INDIAN STOCK HISTORICAL DATA DOWNLOADER"
    )

    print("=" * 75)

    print(
        "Input file:"
    )

    print(
        f"  {INPUT_FILE}"
    )

    print(
        "\nDate range:"
    )

    print(
        f"  {FROM_DATE} -> {TO_DATE}"
    )

    print(
        "\nOutput folder:"
    )

    print(
        f"  {OUTPUT_FOLDER.resolve()}"
    )

    print(
        "\nIndividual files:"
    )

    print(
        "  Yes"
    )

    print(
        "\nMaster workbook:"
    )

    print(
        f"  {MASTER_FILE}"
    )

    print(
        "\nDownload timeout:"
    )

    print(
        f"  {DOWNLOAD_TIMEOUT // 60000} minutes"
    )

    print("=" * 75)

    # --------------------------------------------------------
    # Read tickers
    # --------------------------------------------------------

    tickers = read_tickers()

    # --------------------------------------------------------
    # Stop cleanly if no tickers were found
    # --------------------------------------------------------

    if not tickers:

        print("\n")
        print("=" * 75)
        print(
            "PROCESS STOPPED"
        )
        print("=" * 75)

        print(
            "\nNo tickers were found."
        )

        print(
            "\nPlease check:"
        )

        print(
            "1. CompanyTicker.xlsx exists."
        )

        print(
            "2. Tickers are present in Column A."
        )

        print(
            "3. Column A is not only a header."
        )

        print("\n")

        return

    # --------------------------------------------------------
    # Create output folder
    # --------------------------------------------------------

    OUTPUT_FOLDER.mkdir(
        parents=True,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Create master workbook
    # --------------------------------------------------------

    initialize_master_workbook()

    # --------------------------------------------------------
    # Results log
    # --------------------------------------------------------

    results = []

    # --------------------------------------------------------
    # Start browser
    # --------------------------------------------------------

    with sync_playwright() as p:

        browser = p.chromium.launch(
            headless=HEADLESS
        )

        context = browser.new_context(
            accept_downloads=True
        )

        page = context.new_page()

        # ----------------------------------------------------
        # Process each ticker
        # ----------------------------------------------------

        for number, ticker in enumerate(
            tickers,
            start=1
        ):

            print("\n\n")

            print("=" * 75)

            print(
                f"PROCESSING "
                f"{number}/{len(tickers)}"
            )

            print(
                f"TICKER: {ticker}"
            )

            print("=" * 75)

            exchange = ""
            status = ""
            individual_file = ""
            master_status = ""
            error = ""

            # =================================================
            # FIRST: TRY NSE
            # =================================================

            (
                nse_success,
                nse_file,
                nse_error
            ) = try_nse(
                page,
                ticker
            )

            if nse_success:

                exchange = "NSE"

                status = "Downloaded"

                individual_file = str(
                    nse_file
                )

                # Add to master
                master_success = (
                    add_to_master_workbook(
                        nse_file,
                        ticker,
                        "NSE"
                    )
                )

                if master_success:

                    master_status = "Added"

                else:

                    master_status = "Failed"

                    error = (
                        "Individual download succeeded, "
                        "but adding to master workbook failed."
                    )

            else:

                # =================================================
                # SECOND: TRY BSE
                # =================================================

                print(
                    f"\nNSE failed for {ticker}."
                )

                print(
                    "Trying BSE..."
                )

                (
                    bse_success,
                    bse_file,
                    bse_error
                ) = try_bse(
                    page,
                    ticker
                )

                if bse_success:

                    exchange = "BSE"

                    status = "Downloaded"

                    individual_file = str(
                        bse_file
                    )

                    # Add to master
                    master_success = (
                        add_to_master_workbook(
                            bse_file,
                            ticker,
                            "BSE"
                        )
                    )

                    if master_success:

                        master_status = "Added"

                    else:

                        master_status = "Failed"

                        error = (
                            "Individual download succeeded, "
                            "but adding to master workbook failed."
                        )

                else:

                    exchange = "NSE -> BSE"

                    status = "FAILED"

                    master_status = "Not added"

                    error = (
                        f"NSE error: {nse_error} | "
                        f"BSE error: {bse_error}"
                    )

            # ------------------------------------------------
            # Record result
            # ------------------------------------------------

            results.append({

                "Ticker": ticker,

                "Exchange": exchange,

                "Status": status,

                "Individual File": individual_file,

                "Master Workbook": master_status,

                "Error": error
            })

            # ------------------------------------------------
            # Save log after EVERY company
            # ------------------------------------------------

            save_log(
                results
            )

            # ------------------------------------------------
            # Delay before next company
            # ------------------------------------------------

            if number < len(tickers):

                print(
                    f"\nWaiting "
                    f"{DELAY_BETWEEN_COMPANIES} seconds "
                    f"before next ticker..."
                )

                time.sleep(
                    DELAY_BETWEEN_COMPANIES
                )

        # ----------------------------------------------------
        # Close browser
        # ----------------------------------------------------

        browser.close()

    # ========================================================
    # FINAL SUMMARY
    # ========================================================

    results_df = pd.DataFrame(
        results
    )

    print("\n\n")

    print("=" * 75)

    print(
        "PROCESS COMPLETE"
    )

    print("=" * 75)

    # --------------------------------------------------------
    # Safety check
    # --------------------------------------------------------

    if results_df.empty:

        print(
            "No companies were processed."
        )

        return

    total = len(
        results_df
    )

    successful = len(
        results_df[
            results_df["Status"]
            == "Downloaded"
        ]
    )

    failed = len(
        results_df[
            results_df["Status"]
            == "FAILED"
        ]
    )

    print(
        f"Total tickers : {total}"
    )

    print(
        f"Downloaded    : {successful}"
    )

    print(
        f"Failed        : {failed}"
    )

    print(
        "\nOutput folder:"
    )

    print(
        OUTPUT_FOLDER.resolve()
    )

    print(
        "\nMaster workbook:"
    )

    print(
        MASTER_FILE.resolve()
    )

    print(
        "\nDownload log:"
    )

    print(
        LOG_FILE.resolve()
    )

    print("\n")

    print("=" * 75)


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    main()